#!/usr/bin/env python3
"""
Convertit les exports FileMaker (UF, inventaire préliminaire du mobilier,
objets/prélèvements, inventaire analytique céramique) au format d'import
Excel SIAMOIS, avec un classeur par intervention (= projet), et génère le
thésaurus SKOS des valeurs rencontrées.

Chaque classeur s'importe depuis l'écran d'import du projet correspondant
(portée PROJECT) : le projet doit exister, l'institution et le projet sont
déduits de l'écran, les vocabulaires sont résolus par libellé dans le
thésaurus configuré pour l'institution.

Mapping :
  - chantier                      -> Lieu « Chantier N » (type « Chantier »)
  - UF                            -> UE (type « Unité de fouille »)
  - INventaireprelimobilier       -> Mobilier « Lot »       identifiant <UF>-L<n>
  - objetprelev                   -> Mobilier « Individu » / « Échantillon »
                                     identifiant <UF>-O<numero objet>
  - inventaireanalytiqueceram     -> Mobilier « Lot »       identifiant <UF>-C<n>
  - UF citée par le mobilier mais absente de UF.xlsx -> UE créée (signalée)

Seul le socle commun du modèle d'import est alimenté : année, poids, objet
dateur, sorti, référencier… relèvent des champs additionnels et sont ignorés
(listés dans le rapport).

Sorties (dans --output-dir) :
  - import_intervention_<N>.xlsx  un classeur par intervention
  - thesaurus_filemaker.rdf       SKOS RDF/XML : racine SIAMOIS#SIAAUTO > tables > champs
                                  (SIAMOIS#<code champ>) > valeurs
  - rapport.md                    anomalies et choix appliqués

Usage:
    python3 filemaker_to_siamois.py --input-dir ~/Documents --output-dir ~/Documents/siamois_import

Dépendance externe unique : `openpyxl` (pip install openpyxl).
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
import warnings
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

AUTHOR = "Import FileMaker"

UF_FILE = "UF.xlsx"
INVENTAIRE_FILE = "INventaireprelimobilier.xlsx"
OBJET_FILE = "objetprelev.xlsx"
CERAMIQUE_FILE = "inventaireanalytiqueceram.xlsx"

INTERVENTION = "numero d intervention"


# ---------------------------------------------------------------------------
# Thésaurus : un arbre par champ SIAMOIS. Les racines portent la notation
# SIAMOIS#<code champ> qui permet à SIAMOIS de configurer le champ ; les
# libellés doivent être uniques dans chaque arbre (résolution par libellé).
# ---------------------------------------------------------------------------

@dataclass
class Term:
    label: str
    children: list[Term] = field(default_factory=list)
    note: str | None = None
    code: str | None = None  # code SIAMOIS porté en skos:notation « SIAMOIS#<code> »

    def add(self, label: str, note: str | None = None) -> Term:
        for c in self.children:
            if c.label == label:
                return c
        t = Term(label, note=note)
        self.children.append(t)
        return t

    def walk(self):
        yield self
        for c in self.children:
            yield from c.walk()


LIEU_CHANTIER = "Chantier"
UE_UF = "Unité de fouille"
CAT_LOT, CAT_INDIVIDU, CAT_ECHANTILLON = "Lot", "Individu", "Échantillon"

MATIERE_TREE = [
    Term("Terre cuite"),
    Term("Pierre"),
    Term("Métal", [Term("Fer"), Term("Alliage cuivreux"), Term("Autre métal")]),
    Term("Matière organique", [Term("Os / coquille"), Term("Bois"), Term("Autre matière organique")]),
]

# Les branches « Catégories céramiques » et « Types céramiques » sont
# remplies à partir des valeurs trouvées dans l'inventaire analytique.
DESIGNATION_TREE = [
    Term("Céramique", [
        Term("Vaisselle"),
        Term("Amphore"),
        Term("Catégories céramiques"),
        Term("Types céramiques"),
    ]),
    Term("Élément d'architecture"),
    Term("Quincaillerie", [
        Term("Clou", [Term("Clou à tête plate"), Term("Clou de « murus gallicus »")]),
        Term("Tige"),
        Term("Clé"),
        Term("Anneau"),
        Term("Quincaillerie indéterminée"),
    ]),
    Term("Parure", [Term("Fibule")]),
    Term("Culturel", [Term("Jeton")]),
    Term("Outil", [
        Term("Couteau / lame"),
        Term("Instrument de mouture", [Term("Meule")]),
        Term("Outil indéterminé"),
    ]),
    Term("Prélèvement", [
        Term("Prélèvement botanique"),
        Term("Prélèvement C14"),
        Term("Prélèvement dendrochronologique"),
        Term("Prélèvement autre"),
    ]),
    Term("Scorie"),
    Term("Matériau brut"),
    Term("Autre objet"),
    Term("Indéterminé"),
]

# « categorie mobilier » FileMaker -> (matière, désignation, est un prélèvement).
# Le qualificatif « gardée » (amphore / architecture) n'a pas d'équivalent.
CATEGORIE_MOBILIER = {
    "terre cuite vaisselle":           ("Terre cuite", "Vaisselle", False),
    "terre cuite amphore gardée":      ("Terre cuite", "Amphore", False),
    "terre cuite architecture gardée": ("Terre cuite", "Élément d'architecture", False),
    "terre cuite autre objet":         ("Terre cuite", "Autre objet", False),
    "terre cuite indéterminé":         ("Terre cuite", "Indéterminé", False),
    "fer clous":                       ("Fer", "Clou", False),
    "fer autre objet":                 ("Fer", "Autre objet", False),
    "fer indéterminé":                 ("Fer", "Indéterminé", False),
    "alliage cuivreux autre objet":    ("Alliage cuivreux", "Autre objet", False),
    "alliage cuivreux indéterminé":    ("Alliage cuivreux", "Indéterminé", False),
    "autre métal scorie":              ("Autre métal", "Scorie", False),
    "pierre autre objet":              ("Pierre", "Autre objet", False),
    "pierre autres, indéterminé":      ("Pierre", "Indéterminé", False),
    "pierre instrument de mouture":    ("Pierre", "Instrument de mouture", False),
    "bois autre objet":                ("Bois", "Autre objet", False),
    "bois indéterminé":                ("Bois", "Indéterminé", False),
    "os/coquille brut":                ("Os / coquille", "Matériau brut", False),
    "autre organique brut":            ("Autre matière organique", "Matériau brut", False),
    "prélèvement":                     (None, "Prélèvement", True),
}

# « type mobilier » FileMaker (espaces normalisés) -> (désignation, est un prélèvement).
TYPE_MOBILIER = {
    "Quincaille Clou à tête plate":               ("Clou à tête plate", False),
    'Quincaille Clou de "murus gallicus':         ("Clou de « murus gallicus »", False),
    "Quincaille Tige":                            ("Tige", False),
    "Quincaille Clé":                             ("Clé", False),
    "Quincaille Anneau":                          ("Anneau", False),
    "Quincaille Autre":                           ("Quincaillerie indéterminée", False),
    "Culturel Jeton":                             ("Jeton", False),
    "Parure Fibule":                              ("Fibule", False),
    "Outil Autre":                                ("Outil indéterminé", False),
    "Outil Meule":                                ("Meule", False),
    "Outil Couteau / lame":                       ("Couteau / lame", False),
    "Céramique cf. catégories normalisées":       ("Céramique", False),
    "Prélèvement Botanique":                      ("Prélèvement botanique", True),
    "Prélèvement C14":                            ("Prélèvement C14", True),
    "Prélèvement Dendro":                         ("Prélèvement dendrochronologique", True),
    "Prélèvement Autre":                          ("Prélèvement autre", True),
}

# Libellés céramiques qui entreraient en collision avec un autre terme de l'arbre.
TYPE_CERAMIQUE_RENAME = {"Indéterminée": "Type céramique indéterminé"}

IGNORED_COLUMNS = {
    UF_FILE: ["annee", "LS_couleur stratifiant", "LS_selection stratifiant", "num operation geotopo"],
    INVENTAIRE_FILE: ["annee", "poids"],
    OBJET_FILE: ["annee", "objet dateur", "referencier", "BDB213::sorti"],
    CERAMIQUE_FILE: ["annee", "poids en g", "BDB203 lien 215::concatene uf englobante"],
}


# ---------------------------------------------------------------------------
# Lecture des exports
# ---------------------------------------------------------------------------

def read_rows(path: Path) -> list[dict[str, str | None]]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # exports FileMaker sans style par défaut
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    result = []
    for values in rows:
        row = {h: clean(v) for h, v in zip(header, values)}
        if any(v is not None for v in row.values()):
            result.append(row)
    wb.close()
    return result


def clean(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def squash(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def to_int(s: str | None) -> int | None:
    if s is None:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# Construction des lignes d'import
# ---------------------------------------------------------------------------

@dataclass
class Specimen:
    uf: str
    identifier: str
    category: str
    material: str | None = None
    designation: str | None = None
    count: int | None = None
    tpq: int | None = None


@dataclass
class Project:
    intervention: str
    ufs: dict[str, str] = field(default_factory=dict)            # UF -> chantier
    created_ufs: dict[str, str] = field(default_factory=dict)    # UF -> source ayant forcé la création
    specimens: list[Specimen] = field(default_factory=list)


class Report:
    def __init__(self):
        self.lines: dict[str, list[str]] = defaultdict(list)

    def add(self, section: str, line: str):
        self.lines[section].append(line)


def chantier_name(chantier: str) -> str:
    return f"Chantier {chantier}"


def ensure_uf(project: Project, uf: str, chantier: str | None, source: str, report: Report):
    if uf in project.ufs:
        if chantier and project.ufs[uf] and project.ufs[uf] != chantier:
            report.add("Chantier contradictoire",
                       f"intervention {project.intervention}, UF {uf} : chantier {project.ufs[uf]} "
                       f"dans UF.xlsx, {chantier} dans {source} (conservé : {project.ufs[uf]})")
        return
    project.ufs[uf] = chantier
    project.created_ufs[uf] = source


def build(input_dir: Path, report: Report) -> tuple[dict[str, Project], set[str], set[str]]:
    projects: dict[str, Project] = {}
    ceram_categories: set[str] = set()
    ceram_types: set[str] = set()

    def project(row) -> Project:
        key = row[INTERVENTION]
        return projects.setdefault(key, Project(key))

    for row in read_rows(input_dir / UF_FILE):
        p = project(row)
        if row["UF"] in p.ufs:
            report.add("UF en double", f"intervention {p.intervention}, UF {row['UF']}")
        p.ufs[row["UF"]] = row["chantier"]

    lot_counter: dict[tuple[str, str, str], int] = defaultdict(int)

    def next_number(p: Project, uf: str, prefix: str) -> str:
        lot_counter[(p.intervention, uf, prefix)] += 1
        return f"{uf}-{prefix}{lot_counter[(p.intervention, uf, prefix)]}"

    for row in read_rows(input_dir / INVENTAIRE_FILE):
        p = project(row)
        ensure_uf(p, row["UF"], row["chantier"], INVENTAIRE_FILE, report)
        material, designation, is_sample = map_categorie(row["categorie mobilier"], INVENTAIRE_FILE, report)
        p.specimens.append(Specimen(
            uf=row["UF"],
            identifier=next_number(p, row["UF"], "L"),
            category=CAT_ECHANTILLON if is_sample else CAT_LOT,
            material=material,
            designation=designation,
            count=to_int(row["nombre de fragments"]),
        ))

    for row in read_rows(input_dir / OBJET_FILE):
        p = project(row)
        ensure_uf(p, row["UF"], row["chantier"], OBJET_FILE, report)
        material, designation, is_sample = map_categorie(row["categorie mobilier"], OBJET_FILE, report)
        raw_type = row["type mobilier"]
        if raw_type:
            mapped = TYPE_MOBILIER.get(squash(raw_type))
            if mapped is None:
                report.add("Valeur non mappée", f"{OBJET_FILE} / type mobilier : « {raw_type} »")
            else:
                designation, type_is_sample = mapped
                is_sample = is_sample or type_is_sample
        p.specimens.append(Specimen(
            uf=row["UF"],
            identifier=f"{row['UF']}-O{row['numero objet']}",
            category=CAT_ECHANTILLON if is_sample else CAT_INDIVIDU,
            material=material,
            designation=designation,
        ))

    for row in read_rows(input_dir / CERAMIQUE_FILE):
        p = project(row)
        ensure_uf(p, row["UF"], row["chantier"], CERAMIQUE_FILE, report)
        category = row["categorie ceramique"]
        ctype = row["type ceramique"]
        if category:
            ceram_categories.add(category)
        if ctype:
            ctype = TYPE_CERAMIQUE_RENAME.get(ctype, ctype)
            ceram_types.add(ctype)
        p.specimens.append(Specimen(
            uf=row["UF"],
            identifier=next_number(p, row["UF"], "C"),
            category=CAT_LOT,
            material="Terre cuite",
            designation=ctype or category or "Céramique",
            count=to_int(row["nombre restes"]),
            tpq=to_int(row["TPQ"]),
        ))

    for p in projects.values():
        for uf, source in sorted(p.created_ufs.items(), key=lambda kv: uf_sort_key(kv[0])):
            report.add("UF créées (absentes de UF.xlsx)",
                       f"intervention {p.intervention}, UF {uf} (chantier {p.ufs[uf] or '?'}) — citée dans {source}")
        seen = set()
        for s in p.specimens:
            key = (s.uf, s.identifier)
            if key in seen:
                report.add("Identifiant de mobilier en double",
                           f"intervention {p.intervention}, {s.identifier} (seul le premier sera importé)")
            seen.add(key)

    return projects, ceram_categories, ceram_types


def map_categorie(raw: str | None, source: str, report: Report) -> tuple[str | None, str | None, bool]:
    if raw is None:
        return None, None, False
    mapped = CATEGORIE_MOBILIER.get(raw)
    if mapped is None:
        report.add("Valeur non mappée", f"{source} / categorie mobilier : « {raw} »")
        return None, None, False
    return mapped


def uf_sort_key(uf: str):
    n = to_int(uf)
    return (0, n, "") if n is not None else (1, 0, uf)


# ---------------------------------------------------------------------------
# Écriture des classeurs (noms de feuilles et colonnes canoniques de
# ImportSchema : pas besoin de feuille _meta)
# ---------------------------------------------------------------------------

LIEU_COLUMNS = ["nom", "type label"]
UE_COLUMNS = ["identifiant", "type label", "author email", "unite spatiale"]
PRELEV_COLUMNS = ["identifiant", "categorie label", "matiere label", "designation label",
                  "auteur fiche email", "unite d'enregistrement", "nombre d'elements", "tpq"]


def write_workbook(p: Project, path: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    lieu = wb.create_sheet("Unité spatiale")
    lieu.append(LIEU_COLUMNS)
    for chantier in sorted({c for c in p.ufs.values() if c}, key=uf_sort_key):
        lieu.append([chantier_name(chantier), LIEU_CHANTIER])

    ue = wb.create_sheet("UE")
    ue.append(UE_COLUMNS)
    for uf in sorted(p.ufs, key=uf_sort_key):
        chantier = p.ufs[uf]
        ue.append([as_cell(uf), UE_UF, AUTHOR, chantier_name(chantier) if chantier else None])

    prelev = wb.create_sheet("Prelev")
    prelev.append(PRELEV_COLUMNS)
    for s in sorted(p.specimens, key=lambda s: (uf_sort_key(s.uf), natural_key(s.identifier))):
        prelev.append([s.identifier, s.category, s.material, s.designation, AUTHOR,
                       as_cell(s.uf), s.count, s.tpq])

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 40)
    wb.save(path)


def natural_key(s: str):
    """« 793-L2 » avant « 793-L10 »."""
    return [(0, int(part), "") if part.isdigit() else (1, 0, part) for part in re.split(r"(\d+)", s)]


def as_cell(v: str):
    n = to_int(v)
    return n if n is not None and str(n) == v else v


# ---------------------------------------------------------------------------
# Thésaurus SKOS
# ---------------------------------------------------------------------------

RDF = "http://www.w3.org/1999/02/22-rdf-syntax-ns#"
SKOS = "http://www.w3.org/2004/02/skos/core#"
DCT = "http://purl.org/dc/terms/"
XML_LANG = "{http://www.w3.org/XML/1998/namespace}lang"
BASE_URI = "https://siamois.fr/thesaurus/filemaker/"


def build_thesaurus(ceram_categories: set[str], ceram_types: set[str]) -> list[Term]:
    """Les concepts de champ (un par code SIAMOIS alimenté), avec leurs valeurs."""
    designation = Term("Désignation du mobilier", list(DESIGNATION_TREE), code="SIAS.INTERPRETATION")
    ceramique = designation.children[0]
    categories = next(c for c in ceramique.children if c.label == "Catégories céramiques")
    types = next(c for c in ceramique.children if c.label == "Types céramiques")
    for c in sorted(ceram_categories):
        categories.add(c, note="Code de catégorie céramique FileMaker")
    for t in sorted(ceram_types, key=str.casefold):
        types.add(t, note="Type céramique FileMaker")

    return [
        Term("Type de lieu", [Term(LIEU_CHANTIER)], code="SIASU.TYPE"),
        Term("Type d'unité d'enregistrement", [Term(UE_UF)], code="SIARU.TYPE"),
        Term("Catégorie de mobilier", [Term(CAT_LOT), Term(CAT_INDIVIDU), Term(CAT_ECHANTILLON)], code="SIAS.CAT"),
        Term("Matière", list(MATIERE_TREE), code="SIAS.MATIERE"),
        designation,
    ]


# Concepts de table, comme dans th252 (TBL_lieu = SIAMOIS#SIASU…) : le préfixe
# du code de champ (« SIAS » pour « SIAS.CAT ») désigne sa table.
TABLES = {
    "SIASU": "TBL_lieu",
    "SIARU": "TBL_enregistrement",
    "SIAS": "TBL_prélèvement",
}


def siamois_root(fields: list[Term]) -> Term:
    """Structure attendue par SIAMOIS (ConceptApi.fetchFieldsBranch) : une seule
    racine notée SIAMOIS#SIAAUTO, puis les tables, puis les champs."""
    tables: dict[str, Term] = {}
    for f in fields:
        table_code = f.code.split(".")[0]
        table = tables.setdefault(table_code, Term(TABLES[table_code], code=table_code))
        table.children.append(f)
    return Term("Siamois", list(tables.values()), code="SIAAUTO")


def check_labels(fields: list[Term], used: dict[str, set[str]]):
    """Chaque libellé doit être unique dans son arbre (sinon « ambigu » à l'import)
    et chaque libellé utilisé par les classeurs doit exister dans l'arbre."""
    errors = []
    for root in fields:
        labels = [t.label.casefold() for t in root.walk()]
        dups = {label for label in labels if labels.count(label) > 1}
        if dups:
            errors.append(f"{root.code} : libellés en double {sorted(dups)}")
        missing = {u for u in used.get(root.code, set()) if u.casefold() not in labels}
        if missing:
            errors.append(f"{root.code} : libellés utilisés absents du thésaurus {sorted(missing)}")
    if errors:
        raise SystemExit("Thésaurus incohérent :\n  " + "\n  ".join(errors))


def slug(s: str) -> str:
    s = unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").lower() or "x"


def write_skos(root: Term, path: Path):
    ET.register_namespace("rdf", RDF)
    ET.register_namespace("skos", SKOS)
    ET.register_namespace("dcterms", DCT)
    root_el = ET.Element(f"{{{RDF}}}RDF")
    scheme_uri = BASE_URI + "scheme"

    scheme = ET.SubElement(root_el, f"{{{SKOS}}}ConceptScheme", {f"{{{RDF}}}about": scheme_uri})
    title = ET.SubElement(scheme, f"{{{DCT}}}title", {XML_LANG: "fr"})
    title.text = "SIAMOIS — valeurs issues des exports FileMaker"

    def concept(term: Term, uri: str, broader: str | None):
        el = ET.SubElement(root_el, f"{{{SKOS}}}Concept", {f"{{{RDF}}}about": uri})
        label = ET.SubElement(el, f"{{{SKOS}}}prefLabel", {XML_LANG: "fr"})
        label.text = term.label
        ET.SubElement(el, f"{{{SKOS}}}inScheme", {f"{{{RDF}}}resource": scheme_uri})
        if term.code:
            ET.SubElement(el, f"{{{SKOS}}}notation").text = f"SIAMOIS#{term.code}"
        if broader:
            ET.SubElement(el, f"{{{SKOS}}}broader", {f"{{{RDF}}}resource": broader})
        else:
            ET.SubElement(el, f"{{{SKOS}}}topConceptOf", {f"{{{RDF}}}resource": scheme_uri})
            ET.SubElement(scheme, f"{{{SKOS}}}hasTopConcept", {f"{{{RDF}}}resource": uri})
        if term.note:
            note = ET.SubElement(el, f"{{{SKOS}}}editorialNote", {XML_LANG: "fr"})
            note.text = term.note
        child_uris = [f"{uri}/{slug(c.code or c.label)}" for c in term.children]
        for child_uri in child_uris:
            ET.SubElement(el, f"{{{SKOS}}}narrower", {f"{{{RDF}}}resource": child_uri})
        for child, child_uri in zip(term.children, child_uris):
            concept(child, child_uri, uri)

    concept(root, BASE_URI + slug(root.code), None)

    tree = ET.ElementTree(root_el)
    ET.indent(tree)
    tree.write(path, encoding="utf-8", xml_declaration=True)


# ---------------------------------------------------------------------------
# Rapport
# ---------------------------------------------------------------------------

def write_report(projects: dict[str, Project], report: Report, path: Path):
    out = ["# Conversion FileMaker → SIAMOIS", "", "## Classeurs générés", "",
           "| Intervention | Lieux | UE (dont créées) | Mobilier |", "|---|---|---|---|"]
    for key in sorted(projects, key=uf_sort_key):
        p = projects[key]
        chantiers = sorted({c for c in p.ufs.values() if c}, key=uf_sort_key)
        out.append(f"| {key} | {', '.join(chantiers)} | {len(p.ufs)} ({len(p.created_ufs)}) | {len(p.specimens)} |")
    out += ["", "Auteur de toutes les UE et de tout le mobilier : « " + AUTHOR + " ».", "",
            "## Colonnes ignorées (hors socle commun)", ""]
    for f, cols in IGNORED_COLUMNS.items():
        out.append(f"- `{f}` : " + ", ".join(f"`{c}`" for c in cols))
    out.append("- Qualificatif « gardée » de `categorie mobilier` (amphore / architecture) : non repris.")
    out.append("- Inventaire analytique : la désignation prend le type céramique s'il est renseigné, "
               "sinon la catégorie ; l'autre valeur n'est pas reprise.")
    for section, lines in report.lines.items():
        out += ["", f"## {section} ({len(lines)})", ""]
        out += [f"- {line}" for line in lines]
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input-dir", type=Path, required=True, help="dossier contenant les 4 exports FileMaker")
    parser.add_argument("--output-dir", type=Path, required=True, help="dossier de sortie (créé si besoin)")
    args = parser.parse_args()

    input_dir = args.input_dir.expanduser()
    output_dir = args.output_dir.expanduser()
    for f in (UF_FILE, INVENTAIRE_FILE, OBJET_FILE, CERAMIQUE_FILE):
        if not (input_dir / f).is_file():
            print(f"Fichier introuvable : {input_dir / f}", file=sys.stderr)
            return 1
    output_dir.mkdir(parents=True, exist_ok=True)

    report = Report()
    projects, ceram_categories, ceram_types = build(input_dir, report)

    fields = build_thesaurus(ceram_categories, ceram_types)
    specimens = [s for p in projects.values() for s in p.specimens]
    check_labels(fields, {
        "SIASU.TYPE": {LIEU_CHANTIER},
        "SIARU.TYPE": {UE_UF},
        "SIAS.CAT": {s.category for s in specimens},
        "SIAS.MATIERE": {s.material for s in specimens if s.material},
        "SIAS.INTERPRETATION": {s.designation for s in specimens if s.designation},
    })

    for key, p in projects.items():
        write_workbook(p, output_dir / f"import_intervention_{key}.xlsx")
    write_skos(siamois_root(fields), output_dir / "thesaurus_filemaker.rdf")
    write_report(projects, report, output_dir / "rapport.md")

    print(f"{len(projects)} classeurs, thésaurus et rapport écrits dans {output_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
